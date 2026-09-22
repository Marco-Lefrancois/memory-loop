"""
ingest_file_processors.py - File type parsers for Ingest Agent (ADR-0202 <=300L).

Extracted from ingest_agent.py to comply with modular ceiling.
"""

import hashlib
import json
import re
from pathlib import Path
from datetime import datetime
from collections import Counter
from typing import Optional

from src.cli import ZeroFluffConsole


def compute_sha256(filepath: Path) -> Optional[str]:
    """Compute SHA-256 hash of file content."""
    try:
        file_content_raw = filepath.read_bytes()
        return hashlib.sha256(file_content_raw).hexdigest()
    except Exception as e:
        ZeroFluffConsole.error(f"Impossible de lire le fichier {filepath.name} : {e}")
        return None


def parse_svg(filepath: Path) -> str:
    """Parse SVG file to Markdown text."""
    from src.converters.svg_to_md import parse_svg_to_md_text

    try:
        raw_svg = filepath.read_text(encoding="utf-8", errors="replace")
        return parse_svg_to_md_text(raw_svg, filepath.name)
    except Exception as e:
        return f"[Erreur parsing SVG : {e}]"


def parse_vtt(filepath: Path) -> str:
    """Parse WebVTT file to readable Markdown."""
    try:
        raw = filepath.read_text(encoding="utf-8", errors="replace")
        lines = raw.splitlines()
        dialogue = []
        title = filepath.stem.replace("_", " ")
        dialogue.append(f"# 🎙️ Transcription : {title}\n")

        seen_cues = set()
        for line in lines:
            l = line.strip()
            if not l or l == "WEBVTT" or "-->" in l or l.isdigit():
                continue
            speaker_match = re.match(r"<v\s+([^>]+)>(.*)", l)
            if speaker_match:
                speaker, text = speaker_match.groups()
                text = text.replace("</v>", "").strip()
                cue_key = f"{speaker}:{text}"
                if cue_key not in seen_cues:
                    seen_cues.add(cue_key)
                    dialogue.append(f"- **{speaker}** : {text}")
            else:
                if l not in seen_cues:
                    seen_cues.add(l)
                    dialogue.append(f"- {l}")
        return "\n".join(dialogue)
    except Exception as e:
        return f"[Erreur lors du parsing VTT ({filepath.name}) : {e}]"


def parse_pdf(filepath: Path) -> str:
    """Parse PDF file to text."""
    try:
        import pypdf

        reader = pypdf.PdfReader(filepath)
        return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
    except ImportError:
        msg = "[Erreur : pypdf n'est pas installé. Extraction brute impossible]"
        ZeroFluffConsole.error(f"Dépendance manquante pour PDF ({filepath.name}) : pypdf")
        return msg
    except Exception as e:
        msg = f"[Erreur lors du parsing PDF : {e}]"
        ZeroFluffConsole.error(f"Erreur parsing PDF ({filepath.name}) : {e}")
        return msg


def parse_xlsx(filepath: Path) -> str:
    """Parse XLSX file to Markdown."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheets_text = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            sheets_text.append(f"## Feuille : {sheetname}")
            rows_text = []
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(val) if val is not None else "" for val in row]
                if any(row_vals):
                    rows_text.append(", ".join(row_vals))
            sheets_text.append("\n".join(rows_text))
        return "\n\n".join(sheets_text)
    except ImportError:
        msg = "[Erreur : openpyxl n'est pas installé. Lecture XLSX impossible]"
        ZeroFluffConsole.error(f"Dépendance manquante pour XLSX ({filepath.name}) : openpyxl")
        return msg
    except Exception as e:
        msg = f"[Erreur lors du parsing XLSX : {e}]"
        ZeroFluffConsole.error(f"Erreur parsing XLSX ({filepath.name}) : {e}")
        return msg


def parse_docx(filepath: Path) -> str:
    """Parse DOCX file to text."""
    try:
        import docx

        doc = docx.Document(str(filepath))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        return "\n".join(paragraphs)
    except ImportError:
        msg = "[Erreur : python-docx n'est pas installé. Lecture DOCX impossible]"
        ZeroFluffConsole.error(f"Dépendance manquante pour DOCX ({filepath.name}) : python-docx")
        return msg
    except Exception as e:
        msg = f"[Erreur lors du parsing DOCX : {e}]"
        ZeroFluffConsole.error(f"Erreur parsing DOCX ({filepath.name}) : {e}")
        return msg


def parse_pptx(filepath: Path) -> str:
    """Parse PPTX file to text."""
    try:
        from pptx import Presentation

        prs = Presentation(str(filepath))
        slides_text = []
        for i, slide in enumerate(prs.slides, start=1):
            slides_text.append(f"## Slide {i}")
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text:
                    slides_text.append(text)
        return "\n\n".join(slides_text)
    except ImportError:
        msg = "[Erreur : python-pptx n'est pas installé. Lecture PPTX impossible]"
        ZeroFluffConsole.error(f"Dépendance manquante pour PPTX ({filepath.name}) : python-pptx")
        return msg
    except Exception as e:
        msg = f"[Erreur lors du parsing PPTX : {e}]"
        ZeroFluffConsole.error(f"Erreur parsing PPTX ({filepath.name}) : {e}")
        return msg


def parse_csv(filepath: Path) -> str:
    """Parse CSV file to Markdown summary."""
    try:
        from src.converters.csv_engine import csv_to_markdown_summary

        return csv_to_markdown_summary(filepath, max_preview_rows=25)
    except Exception as e:
        return f"[Erreur de lecture CSV : {e}]"


def try_markitdown(filepath: Path) -> str:
    """Universal ingestion via Microsoft MarkItDown."""
    try:
        from markitdown import MarkItDown

        md = MarkItDown()
        res = md.convert(str(filepath))
        return res.text_content if hasattr(res, "text_content") else str(res)
    except Exception:
        return ""


def parse_text_file(filepath: Path) -> str:
    """Parse plain text/markdown/cs files with encoding fallback."""
    try:
        return filepath.read_text(encoding="utf-8")
    except Exception:
        try:
            return filepath.read_text(encoding="cp1252")
        except Exception as e:
            return f"[Erreur de lecture texte : {e}]"


def process_image_asset(filepath: Path, state) -> str:
    """Sync image to docs/05-assets/ and generate artifact note."""
    try:
        import shutil
        from src.state import ProjectLayout

        rel_parent = filepath.parent.name
        target_assets = Path("Projects") / "mLoop" / "docs" / "05-assets"
        if filepath.parent.name in ["01-reception", "02-incubation", "03-ventes"]:
            target_assets = target_assets / filepath.parent.name
        target_assets.mkdir(parents=True, exist_ok=True)
        target_img = target_assets / filepath.name
        import shutil

        shutil.copy2(filepath, target_img)

        return f"# 🖼️ Capture / Maquette : {filepath.stem}\n\n![{filepath.stem}](../05-assets/{filepath.parent.name}/{filepath.name})\n\n- **Fichier source** : `{filepath.name}`\n- **Module associé** : `{filepath.parent.name}`\n"
    except Exception as e:
        return f"[Erreur lors du traitement de l'image ({filepath.name}) : {e}]"
